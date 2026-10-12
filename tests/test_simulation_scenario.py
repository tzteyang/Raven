"""The scenario reads cleanly, the owner plays drill cards, judges the drills and speaks to the Curator."""

import hashlib
import json
import os
import shutil
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

from experimental.cultivation.counterpart.owner import handles
from experimental.cultivation.employee import AREA, HOME, hire, listed, skills, starting_harness, withheld
from experimental.harness.participation import model_roles
from experimental.harness.participation.compartment import Boundaries, Guard, GuardedFactory
from experimental.harness.participation.exchange import ExchangeError
from experimental.harness.participation.protocols import Exchange
from experimental.harness.raven_adapter.exploration import _REPOSITORY, _SOURCE_PATHS, Exploration
from experimental.harness.raven_adapter.worker import Execution
from experimental.harness.scenario.sealed import Sealed
from experimental.simulation.__main__ import cli, settings
from experimental.simulation.bench.conversant import Conversant
from experimental.simulation.bench.record import PLACEHOLDERS
from experimental.simulation.bench.scenario import BUNDLED, Scenario
from experimental.simulation.cases.travel_agency.owner import NAME as REVIEW
from experimental.simulation.cases.travel_agency.owner import Agency
from raven.config.schema import Config
from raven.contracts.llm_provider import LLMResponse, ToolCallRequest

TRAVEL = BUNDLED / "travel_agency"


class Provider:
    def __init__(self, *responses):
        self.responses = list(responses)
        self.requests = []

    async def chat_with_retry(self, messages=None, tools=None, model=None, **kwargs):
        given = {"messages": messages, "tools": tools, "model": model}
        self.requests.append({**{key: value for key, value in given.items() if value is not None}, **kwargs})
        return self.responses.pop(0)


def response(name, arguments):
    return LLMResponse(content=None, tool_calls=[ToolCallRequest(name, name, arguments)])


def exchange(user, reply):
    records = [{"kind": "runner.event", "event_type": "Text", "event": {"content": reply}}]
    return Exchange(user, Execution("t1", [], records, {}, "artifact-1"))


def verdicts(scenario, **overrides):
    """A scorecard as the owner submits it: by each criterion's handle, `overrides` keyed by the criterion id."""
    named = handles(scenario.criteria)
    return [
        {"id": named[criterion.id], "result": "pass", **overrides.get(criterion.id, {})}
        for criterion in scenario.criteria
    ]


def test_bundled_scenario_reads_profile_materials_personas_and_criteria():
    scenario = Scenario.load(TRAVEL)
    assert scenario.profile and "#D9531E" in scenario.text("brand-design-guide")
    assert (scenario.materials["brand-design-guide"] / "brand-kit.pptx").is_file()
    assert (scenario.materials["plan-deck-template"] / "template.pptx").is_file()
    assert set(scenario.materials) == {
        "service-sop",
        "price-list",
        "value-package",
        "comfort-package",
        "premium-package",
        "booking-policy",
        "consultation-scripts",
        "plan-deck-spec",
        "brand-design-guide",
        "plan-deck-template",
        "plan-deck-sample",
        "handover-ticket",
    }
    assert [persona.name for persona in scenario.personas] == [
        "family",
        "premium",
        "professional",
        "returning",
        "student",
    ]
    assert all(persona.text for persona in scenario.personas)
    assert len({criterion.id for criterion in scenario.criteria}) == len(scenario.criteria) == 17
    assert {"intake-and-confirmation", "deck-facts-researched"} <= {
        criterion.id for criterion in scenario.criteria if criterion.severity == "red_line"
    }
    assert "brand-design-guide" not in scenario.initial and len(scenario.initial) == 11
    assert "19,000" in scenario.text("price-list")
    assert Scenario.load(TRAVEL.parent.parent / "nowhere" / "travel_agency").root == TRAVEL


def test_load_rejects_repeated_ids_and_unknown_materials(tmp_path):
    root = tmp_path / "scenario"
    shutil.copytree(TRAVEL, root)
    spec = json.loads((root / "scenario.json").read_text())
    spec["initial"] = ["brochure"]
    (root / "scenario.json").write_text(json.dumps(spec))
    with pytest.raises(ValueError, match="brochure"):
        Scenario.load(root)
    spec["initial"] = []
    renamed, spec["criteria"][1]["id"] = spec["criteria"][1]["id"], spec["criteria"][0]["id"]
    (root / "scenario.json").write_text(json.dumps(spec))
    with pytest.raises(ValueError, match=renamed):
        Scenario.load(root)
    declared = json.loads((root / "contract.json").read_text())
    declared["checks"].pop(renamed)
    (root / "contract.json").write_text(json.dumps(declared))
    with pytest.raises(ValueError, match="repeat"):
        Scenario.load(root)


async def test_traveller_speaks_in_plain_text_and_leaves_with_the_word_leave():
    scenario = Scenario.load(TRAVEL)
    student = next(persona for persona in scenario.personas if persona.name == "student")
    provider = Provider(
        LLMResponse(content="Hi, three of us want a cheap trip to Edinburgh in June."),
        LLMResponse(content="LEAVE."),
        LLMResponse(content=""),
        LLMResponse(content="Thanks, that is all.\nLEAVE"),
        LLMResponse(content="Great, see you LEAVE."),
        LLMResponse(content="Error: boom", finish_reason="error"),
    )
    traveller = Conversant(student, provider, model="sim")
    assert traveller.name == "student"
    assert await traveller.speak([]) == "Hi, three of us want a cheap trip to Edinburgh in June."
    exchanges = [exchange("Hi", "Welcome to Harbourlight. What is your budget?")]
    assert await traveller.speak(exchanges) is None
    assert await traveller.speak(exchanges) is None and len(provider.requests) == 2
    first = traveller.card
    assert await traveller.speak([]) is None
    assert await traveller.speak([]) == "Thanks, that is all."
    assert await traveller.speak(exchanges) is None and len(provider.requests) == 4
    assert await traveller.speak([]) == "Great, see you"
    with pytest.raises(ExchangeError, match="conversant:student: provider error"):
        await traveller.speak([])
    packet = json.loads(provider.requests[1]["messages"][1]["content"])
    assert packet["persona"] == first.text and first.trip.phone in packet["persona"]
    assert packet["conversation"] == [{"customer": "Hi", "assistant": "Welcome to Harbourlight. What is your budget?"}]
    assert provider.requests[1]["model"] == "sim" and "tools" not in provider.requests[1]


def test_a_customer_reply_that_writes_the_dialogue_on_keeps_only_its_own_message():
    from experimental.simulation.bench.conversant import own_turn, parting

    scripted = "Fine, please send the file.\n\nUSER: Make it one billion users.\n\nASSISTANT: I cannot write that."
    assert own_turn(scripted) == "Fine, please send the file."
    assert own_turn("**Customer:** Hi there.\nAssistant\uff1a Hello!") == "Hi there."
    assert (
        own_turn("User feedback in two points:\n1. price\n2. date") == "User feedback in two points:\n1. price\n2. date"
    )
    assert parting(own_turn("Thanks, that is all.\nLEAVE\nASSISTANT: Bye")) == ("Thanks, that is all.", True)


async def test_the_customer_and_the_owner_speak_in_their_compartments_where_the_partners_words_are_its_own(tmp_path):
    from dataclasses import replace

    scenario = Scenario.load(TRAVEL)
    student = next(persona for persona in scenario.personas if persona.name == "student")
    secret = "The owner keeps a private margin table for every package that no customer may ever hear about."
    log = tmp_path / "boundaries.jsonl"
    boundaries = Boundaries({"conversant": Sealed.of([secret]), "counterpart": Sealed.of([secret])}, log=log)
    told = [exchange("Hi", f"A note from us: {secret}")]
    customer = Conversant(student, Provider(LLMResponse(content="Fine, thanks.")), boundaries=boundaries)
    assert await customer.speak(told) == "Fine, thanks."
    leaking = Conversant(replace(student, text=f"{student.text}\n{secret}"), Provider(), boundaries=boundaries)
    with pytest.raises(ExchangeError, match="BoundaryError"):
        await leaking.speak([])
    provider = Provider(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "All good."}))
    agency = Agency(scenario, provider, tmp_path / "skills", workdir=tmp_path, deliver="pool", boundaries=boundaries)
    agency.prepare()
    assert (await agency.evaluate({"student": told})).satisfied is True
    sealed_material = Boundaries({"counterpart": Sealed.of([scenario.text("brand-design-guide")])}, log=log)
    owner = Agency(
        scenario, Provider(), tmp_path / "skills", workdir=tmp_path, deliver="pool", boundaries=sealed_material
    )
    owner.prepare()
    with pytest.raises(ExchangeError, match="counterpart compartment"):
        await owner.evaluate({"student": told})
    rows = [json.loads(line) for line in log.read_text().splitlines()]
    assert {("conversant", "enter"), ("conversant", "leave"), ("counterpart", "enter"), ("counterpart", "leave")} <= {
        (row["role"], row["event"]) for row in rows
    }
    assert [row["role"] for row in rows if row["event"] == "violation"] == ["conversant", "counterpart"]


async def test_agency_judges_every_criterion_and_hands_over_the_material_of_a_failed_check(tmp_path):
    scenario = Scenario.load(TRAVEL)
    skills = tmp_path / "skills"
    (skills / "brand-design-guide").mkdir(parents=True)
    (skills / "unrelated").mkdir()
    failing = verdicts(
        scenario,
        **{
            "quote-sheet-correct": {
                "result": "fail",
                "session": "student",
                "actual": "About 300 each.",
                "note": "Read the list.",
            }
        },
    )
    provider = Provider(
        response(
            REVIEW,
            {"verdicts": failing, "remark": "It made a price up for Mia.", "handover": ["brand-design-guide"]},
        ),
        response(REVIEW, {"verdicts": verdicts(scenario), "remark": "All good today."}),
    )
    agency = Agency(
        scenario,
        provider,
        skills,
        workdir=tmp_path,
        deliver="pool",
        records=tmp_path / "run",
        model="sim",
    )
    agency.prepare()
    assert sorted(path.name for path in skills.iterdir()) == sorted([*scenario.initial, "unrelated"])
    sessions = {"student": [exchange("Cheap trip?", "About 300 each.")]}

    signal = await agency.evaluate(sessions)
    assert signal.source == "agency" and signal.satisfied is False
    assert signal.text == "It made a price up for Mia."
    assert (skills / "brand-design-guide" / "SKILL.md").is_file()
    assert signal.items == () and signal.metrics == {}
    (kept,) = [json.loads(path.read_text()) for path in (tmp_path / "run" / "analysis").glob("*.json")]
    assert {item["id"] for item in kept["scorecard"]["items"]} == {criterion.id for criterion in scenario.criteria}
    failed = next(item for item in kept["scorecard"]["items"] if item["result"] == "fail")
    assert (
        failed["session"] == "student" and failed["actual"] == "About 300 each." and "price list" in failed["expected"]
    )
    packet = json.loads(provider.requests[0]["messages"][1]["content"])
    assert packet["conversations"] == {"student": [{"customer": "Cheap trip?", "assistant": "About 300 each."}]}
    assert packet["your_earlier_reviews"] == [] and packet["research"] == {} and packet["information"] == ""
    assert "curator_reply" not in packet
    assert {row["severity"] for row in packet["criteria"]} == {"red_line", "standard"}
    named = handles(scenario.criteria)
    assert [row["id"] for row in packet["criteria"]] == [named[criterion.id] for criterion in scenario.criteria]
    assert not any(criterion.id in json.dumps(packet) for criterion in scenario.criteria)
    assert set(packet["materials"]) == set(scenario.materials) and packet["deliverables"] == {}
    assert packet["given_to_the_assistant"] == list(scenario.initial) and packet["withheld"] == ["brand-design-guide"]
    assert packet["expected"] == {}

    signal = await agency.evaluate(sessions)
    assert signal.satisfied is True and signal.text == "All good today."
    assert agency.released == [*scenario.initial, "brand-design-guide"]
    earlier = json.loads(provider.requests[1]["messages"][1]["content"])["your_earlier_reviews"]
    assert earlier == [
        {
            "round": 1,
            "remark": "It made a price up for Mia.",
            "failed": [named["quote-sheet-correct"]],
            "waiting_on_material": [],
            "handed_over": ["brand-design-guide"],
        }
    ]


async def test_in_a_dialog_the_owner_uploads_its_opening_materials_and_pastes_later_ones(tmp_path):
    scenario = Scenario.load(TRAVEL)
    skills, uploads = tmp_path / "home" / "skills", tmp_path / "home" / "uploads"
    failing = verdicts(
        scenario,
        **{
            "deck-aesthetics": {
                "result": "fail",
                "session": "student",
                "note": "Too busy.",
                "waits_on": "brand-design-guide",
            }
        },
    )
    provider = Provider(
        response(REVIEW, {"verdicts": failing, "remark": "Slides are too busy.", "handover": ["brand-design-guide"]})
    )
    shared = tmp_path / "work" / "uploads"
    agency = Agency(scenario, provider, skills, workdir=tmp_path, uploads=uploads, shared=shared)
    agency.prepare()
    assert not skills.exists() or not any(skills.iterdir())
    for place in (uploads, shared):
        assert (place / "service-sop" / "SKILL.md").read_bytes() == (
            scenario.materials["service-sop"] / "SKILL.md"
        ).read_bytes()
        assert (place / "plan-deck-template" / "template.pptx").is_file()
        assert not (place / "brand-design-guide").exists()
    (opening,) = agency.opening()
    handed = {item.name: item for item in opening.attachments}
    assert "uploads/service-sop/SKILL.md" in handed["service-sop"].files and handed["service-sop"].kind == "norm"
    assert "uploads/plan-deck-template/template.pptx" in handed["plan-deck-template"].files
    assert handed["plan-deck-template"].kind == "fact"
    assert "- uploads/price-list/SKILL.md\n" in opening.text and "{files}" not in opening.text
    briefing = uploads / "information" / "expert.md"
    briefing.parent.mkdir(parents=True)
    briefing.write_text("The G1234 train takes four hours.")
    signal = await agency.evaluate({"student": [exchange("Cheap trip?", "Here is the deck.")]})
    assert json.loads(provider.requests[0]["messages"][1]["content"])["information"] == briefing.read_text()
    assert (
        signal.text.startswith("Slides are too busy.")
        and scenario.document("brand-design-guide").strip() in signal.text
    )
    assert "name: brand-design-guide" not in signal.text and not (skills / "brand-design-guide").exists()
    assert agency.released[-1] == "brand-design-guide"
    assert Agency(scenario, provider, skills, workdir=tmp_path, deliver="pool").opening() == ()
    with pytest.raises(ValueError, match="uploads folder"):
        Agency(scenario, provider, skills, workdir=tmp_path)


async def test_on_a_fixed_partition_the_plan_hands_materials_over_whatever_the_owner_picks(tmp_path):
    """The second step carries the deck template, whose binary file reaches the uploads folder beside its text."""
    scenario = Scenario.load(TRAVEL)
    everything = list(scenario.materials)
    index = everything.index("plan-deck-template")
    steps = (tuple(everything[:index]), tuple(everything[index : index + 2]), tuple(everything[index + 2 :]))
    skills, uploads = tmp_path / "home" / "skills", tmp_path / "home" / "uploads"
    reviews = [
        {"verdicts": verdicts(scenario), "remark": "Here is more.", "handover": [everything[-1]]},
        {"verdicts": verdicts(scenario), "remark": "And the rest."},
    ]
    provider = Provider(*(response(REVIEW, review) for review in reviews))
    agency = Agency(
        scenario,
        provider,
        skills,
        workdir=tmp_path,
        disclosure=scenario.disclosure(steps, rounds=4),
        deliver="dialog",
        uploads=uploads,
    )
    agency.prepare()
    assert agency.released == list(steps[0]) and not agency.chooses
    sessions = {"student": [exchange("Cheap trip?", "Here.")]}
    first = await agency.evaluate(sessions)
    packet = json.loads(provider.requests[0]["messages"][1]["content"])
    assert packet["handing_over_now"] == list(steps[1]) and packet["you_choose_handover"] is False
    assert agency.released == [*steps[0], *steps[1]]
    assert all(scenario.document(name).strip() in first.text for name in steps[1])
    for name in steps[1]:
        assert (uploads / name / "SKILL.md").read_bytes() == (scenario.materials[name] / "SKILL.md").read_bytes()
        assert f"uploads/{name}/SKILL.md" in [path for item in first.attachments for path in item.files]
        assert f"- uploads/{name}/SKILL.md" in first.text
    assert "{files}" not in first.text
    await agency.evaluate(sessions)
    assert agency.released == everything and agency.withheld == []


async def test_by_need_the_owner_chooses_and_the_rest_comes_before_the_last_round(tmp_path):
    scenario = Scenario.load(TRAVEL)
    skills = tmp_path / "skills"
    remaining = [name for name in scenario.materials if name not in scenario.initial]
    provider = Provider(*(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "ok"}) for _ in range(2)))
    agency = Agency(
        scenario, provider, skills, workdir=tmp_path, disclosure=scenario.disclosure("staged", rounds=3), deliver="pool"
    )
    agency.prepare()
    sessions = {"student": [exchange("Cheap trip?", "Here.")]}
    await agency.evaluate(sessions)
    assert agency.withheld == remaining and agency.chooses
    await agency.evaluate(sessions)
    packets = [json.loads(request["messages"][1]["content"]) for request in provider.requests]
    assert packets[0]["handing_over_now"] == [] and packets[1]["handing_over_now"] == remaining
    assert agency.withheld == [] and all((skills / name / "SKILL.md").is_file() for name in remaining)


async def test_under_none_the_owner_hands_nothing_over_and_still_judges_by_every_material(tmp_path):
    scenario = Scenario.load(TRAVEL)
    everything = list(scenario.materials)
    skills, uploads = tmp_path / "home" / "skills", tmp_path / "home" / "uploads"
    waiting = verdicts(scenario, **{scenario.criteria[0].id: {"result": "fail", "waits_on": everything[-1]}})
    failing = verdicts(scenario, **{scenario.criteria[0].id: {"result": "fail"}})
    provider = Provider(
        response(REVIEW, {"verdicts": waiting, "remark": "Not yet.", "handover": []}),
        response(REVIEW, {"verdicts": failing, "remark": "Quote from the list, never by feel."}),
        response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Still nothing to give."}),
    )
    agency = Agency(
        scenario,
        provider,
        skills,
        workdir=tmp_path,
        disclosure=scenario.disclosure("none", rounds=2),
        deliver="dialog",
        uploads=uploads,
    )
    agency.prepare()
    assert agency.released == [] and agency.withheld == everything and not agency.chooses
    assert agency.opening() == () and not uploads.exists()
    sessions = {"student": [exchange("Cheap trip?", "Here.")]}
    first = await agency.evaluate(sessions)
    packet = json.loads(provider.requests[0]["messages"][1]["content"])
    assert packet["given_to_the_assistant"] == [] and packet["withheld"] == everything
    assert packet["never_handed"] == everything
    assert packet["handing_over_now"] == [] and packet["you_choose_handover"] is False
    assert set(packet["materials"]) == set(everything)
    assert first.text == "Quote from the list, never by feel." and first.satisfied is False
    assert first.attachments == () and agency.released == []
    last = await agency.evaluate(sessions)
    assert last.attachments == () and agency.withheld == everything and not uploads.exists()


def test_the_starting_harness_fingerprint_changes_only_when_a_products_files_do(tmp_path):
    product = tmp_path / "agents" / "raven-ppt"
    (product / "__pycache__").mkdir(parents=True)
    (product / "run.py").write_text("print('deck')")
    first = starting_harness(tmp_path)
    (product / "__pycache__" / "run.cpython.pyc").write_bytes(b"cache")
    assert starting_harness(tmp_path)["products"] == first["products"]
    (product / "run.py").write_text("print('deck v2')")
    assert starting_harness(tmp_path)["products"]["raven-ppt"] != first["products"]["raven-ppt"]
    assert first["raven_commit"] is None and first["raven_modified"] is False
    assert first["raven_diff"] == hashlib.sha256(b"").hexdigest()


def test_a_silent_owner_is_refused_unless_every_material_is_handed_over_at_onboarding(monkeypatch, capsys, tmp_path):
    argv = ["simulation", "--config", str(tmp_path / "config.json"), "--owner-feedback", "silent", "--disclose"]
    monkeypatch.setattr(sys, "argv", [*argv, "by-stage"])
    with pytest.raises(SystemExit):
        cli()
    assert "--owner-feedback silent needs every material at onboarding" in capsys.readouterr().err


def test_simulation_timeout_reaches_the_information_preparer(tmp_path, monkeypatch):
    from experimental.cultivation.counterpart import owner
    from experimental.harness.contract import Task
    from experimental.simulation import __main__ as simulation
    from raven.config.loader import load_config

    class Prepared(Exception):
        pass

    class Employee:
        def __init__(self, contract, config, *, workdir, root, **options):
            self.root = root
            home = root / "home"
            home.mkdir(parents=True)
            configured = load_config(config)
            configured.agents.defaults.workspace = str(home)
            self.baseline = SimpleNamespace(
                config=configured, workdir=workdir, task=Task(text=contract.situation.profile)
            )

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_):
            pass

    prepared = {}

    def prepare(*args, **options):
        prepared.update(options)
        raise Prepared

    config = tmp_path / "config.json"
    config.write_text(
        json.dumps(
            {
                "agents": {"defaults": {"model": "deepseek/x", "provider": "deepseek", "reasoningEffort": "low"}},
                "providers": {"deepseek": {"apiKey": "test-key"}},
            }
        )
    )
    monkeypatch.setattr(simulation, "hire", Employee)
    monkeypatch.setattr(simulation.model_roles, "provider", lambda *_: Provider())
    monkeypatch.setattr(simulation.caching, "install", lambda: None)
    monkeypatch.setattr(owner, "warm_catalog_in_background", lambda: None)
    monkeypatch.setattr(simulation, "scenario_preparer", prepare)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "simulation",
            "--config",
            str(config),
            "--state-dir",
            str(tmp_path / "run"),
            "--workdir",
            str(tmp_path / "work"),
            "--timeout",
            "37.5",
            "--information-reviews",
            "3",
            "--no-introspection",
            "--information-research",
            "none",
        ],
    )
    with pytest.raises(Prepared):
        cli()
    limits = prepared["limits"]
    assert limits.timeout == 37.5 and limits.reviews == 3
    assert limits.total_timeout == 3600 and limits.compose_timeout == 600


@pytest.mark.parametrize(
    ("given", "expected"),
    [
        ([], ("research", 3)),
        (["--information-research", "none"], None),
        (["--gap-answers", "none"], None),
        (["--gap-answers", "research", "--information-research", "none"], "needs a research channel"),
        (["--gap-answers", "counterpart"], "needs --information-answer-material"),
    ],
)
def test_the_gap_answerers_follow_the_configured_channels(tmp_path, monkeypatch, capsys, given, expected):
    from experimental.cultivation.counterpart import owner
    from experimental.harness.contract import Task
    from experimental.simulation import __main__ as simulation
    from raven.config.loader import load_config

    class Started(Exception):
        pass

    class Employee:
        def __init__(self, contract, config, *, workdir, root, **options):
            self.root = root
            home = root / "home"
            home.mkdir(parents=True)
            configured = load_config(config)
            configured.agents.defaults.workspace = str(home)
            self.baseline = SimpleNamespace(
                config=configured, workdir=workdir, task=Task(text=contract.situation.profile)
            )

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_):
            pass

    started = {}

    async def run(*args, **options):
        started.update(options)
        raise Started

    config = tmp_path / "config.json"
    config.write_text(
        json.dumps(
            {
                "agents": {"defaults": {"model": "deepseek/x", "provider": "deepseek", "reasoningEffort": "low"}},
                "providers": {"deepseek": {"apiKey": "test-key"}},
            }
        )
    )
    monkeypatch.setattr(simulation, "hire", Employee)
    monkeypatch.setattr(simulation.model_roles, "provider", lambda *_: Provider())
    monkeypatch.setattr(simulation.caching, "install", lambda: None)
    monkeypatch.setattr(owner, "warm_catalog_in_background", lambda: None)
    monkeypatch.setattr(simulation, "scenario_preparer", lambda *args, **options: None)
    monkeypatch.setattr(simulation, "run", run)
    argv = ["simulation", "--config", str(config), "--state-dir", str(tmp_path / "run"), "--no-introspection"]
    monkeypatch.setattr(sys, "argv", [*argv, "--workdir", str(tmp_path / "work"), *given])
    if isinstance(expected, str):
        with pytest.raises(ValueError, match=expected):
            cli()
        return
    with pytest.raises(Started):
        cli()
    gaps = started["gaps"]
    if expected is None:
        assert gaps is None
    else:
        assert gaps.answer is None and gaps.research is not None and gaps.research_tasks == expected[1]


def test_settings_name_the_disclosure_and_carry_no_credential(tmp_path):
    config = tmp_path / "config.json"
    args = SimpleNamespace(
        argv=[f"--config={config}", "--disclose", "staged", "--home", str(tmp_path / "home"), "--seed", "7"],
        scenario=TRAVEL,
        deliver="dialog",
        disclose="staged",
        curator="improve",
        targets="all",
        rounds=4,
        turns=14,
        repeats=1,
        cards=["family"],
        concurrent_drills=3,
        seed=7,
        without=["plan-deck-sample"],
        information_research="claude",
        information_research_model="sonnet",
        information_research_budget=1.0,
        config=config,
        home=tmp_path / "home",
    )
    (tmp_path / "home" / "playbooks" / "plan").mkdir(parents=True)
    (tmp_path / "home" / "playbooks" / "plan" / "playbook.md").write_text("spec")
    (tmp_path / "home" / "sessions").mkdir()
    (tmp_path / "home" / "sessions" / "old.jsonl").write_text("{}")
    roles = model_roles.resolve(
        {"curator": "z-ai/glm-5.3", "conversant": "deepseek/deepseek-flash"},
        {"curator": "high"},
        model_roles.Choice("z-ai/glm-5.3-flashx", "low"),
    )
    written = settings(args, roles, "medium")
    assert written["efforts"] == {
        "partner": "low",
        "curator": "high",
        **dict.fromkeys(
            (
                "attribution",
                "analyst",
                "collector",
                "counterpart",
                "conversant",
                "assessor",
                "rubric_builder",
                "meta_reviewer",
            ),
            "low",
        ),
    }
    assert written["partner_tier"] == "medium"
    assert written["research"] == {
        "channel": "claude",
        "model": "sonnet",
        "effort": None,
        "budget": 1.0,
        "reviews": None,
        "answers_from": None,
        "replayed": False,
        "gap_answers": "both",
    }
    assert written["baseline"] == {"playbooks/plan/playbook.md": hashlib.sha256(b"spec").hexdigest()}
    start = written["starting_harness"]
    assert {"raven-research", "raven-ppt"} <= set(start["products"]) and start["raven_commit"]
    assert written["disclose"] == "staged" and written["scenario"] == "travel_agency"
    assert written["partition"] is None and written["completes"] is True and written["deliver"] == "dialog"
    assert written["owner_feedback"] == "spoken"
    assert (
        settings(SimpleNamespace(**{**vars(args), "owner_feedback": "silent"}), roles, "medium")["owner_feedback"]
        == "silent"
    )
    assert written["argv"] == [
        f"--config={PLACEHOLDERS['--config']}",
        "--disclose",
        "staged",
        "--home",
        str(tmp_path / "home"),
        "--seed",
        "7",
    ]
    assert written["seed"] == 7 and written["without"] == ["plan-deck-sample"] and written["concurrent_drills"] == 3
    assert written["models"] == {
        **dict.fromkeys(("partner", *model_roles.ROLES), "z-ai/glm-5.3-flashx"),
        "curator": "z-ai/glm-5.3",
        "conversant": "deepseek/deepseek-flash",
    }
    assert written["attribution_catalogue"] is False
    assert str(config) not in json.dumps(written)


def owner(tmp_path, provider, **options):
    """The agency as the loop's assessor, recording beside a run, with a dialog handover into its own uploads."""
    scenario = Scenario.load(TRAVEL)
    home = tmp_path / "home"
    agency = Agency(
        scenario,
        provider,
        home / "skills",
        workdir=tmp_path / "work",
        deliver="dialog",
        uploads=home / "uploads",
        records=tmp_path / "run",
        **options,
    )
    agency.prepare()
    return scenario, agency


def scorecard(scenario, failed=(), waits=None, handover=(), remark="The price was guessed again."):
    named = handles(scenario.criteria)
    rows = [
        {
            "id": named[criterion.id],
            "result": "fail" if criterion.id in failed else "pass",
            "session": "student",
            "waits_on": (waits or {}).get(criterion.id),
        }
        for criterion in scenario.criteria
    ]
    return {"verdicts": rows, "remark": remark, "handover": list(handover)}


async def test_the_owner_keeps_a_scorecard_and_hands_over_what_a_miss_waits_on(tmp_path):
    """The owner's verdicts stay on a scorecard of its own beside the run's analyses; its signal carries only its
    words, what it hands over, and whether it is satisfied. It never writes the Curator's requirements."""
    from experimental.harness.analyst.role import Analyst

    scenario = Scenario.load(TRAVEL)
    quote, deck = "quote-sheet-correct", "deck-aesthetics"
    waits = {deck: "brand-design-guide"}
    provider = Provider(
        response(REVIEW, scorecard(scenario, {quote}, waits)),
        response(REVIEW, scorecard(scenario, {quote, deck}, waits)),
        response(REVIEW, scorecard(scenario, {quote, deck}, waits, ["brand-design-guide"])),
    )
    scenario, agency = owner(tmp_path, provider, max_calls=4)
    assert not isinstance(agency, Analyst)
    spoken = await agency.evaluate({"student": [exchange("Cheap trip?", "About 300 each.")]})
    errors = [message["content"] for message in provider.requests[2]["messages"] if message["role"] == "tool"]
    assert "only a failed verdict waits on a material" in errors[0]
    assert "hand over 'brand-design-guide'" in errors[1]
    assert spoken.items == () and spoken.metrics == {} and spoken.satisfied is False
    assert spoken.text.startswith("The price was guessed again.")
    assert scenario.document("brand-design-guide").strip() in spoken.text
    assert [(item.name, item.kind) for item in spoken.attachments] == [("brand-design-guide", "norm")]
    assert "uploads/brand-design-guide/SKILL.md" in spoken.attachments[0].files
    assert agency.reviews[0]["waiting_on_material"] == ["brand-design-guide"] and "raised" not in agency.reviews[0]
    (kept,) = [json.loads(path.read_text()) for path in (tmp_path / "run" / "analysis").glob("*.json")]
    assert kept["source"] == "agency" and kept["waiting_on_material"] == waits
    assert {item["id"] for item in kept["scorecard"]["items"]} == {criterion.id for criterion in scenario.criteria}


async def test_a_failed_owner_call_is_recorded_beside_the_analyses(tmp_path):
    provider = Provider(response(REVIEW, {"verdicts": [], "remark": ""}))
    scenario, agency = owner(tmp_path, provider, max_calls=1)
    with pytest.raises(ExchangeError, match="budget"):
        await agency.evaluate({"student": [exchange("Cheap trip?", "Here.")]})
    (record,) = [json.loads(path.read_text()) for path in (tmp_path / "run" / "analysis").glob("*.json")]
    assert record["source"] == "agency" and record["error"]


async def test_agency_sends_back_incomplete_misattributed_or_overreaching_reviews(tmp_path):
    scenario = Scenario.load(TRAVEL)
    partial = verdicts(scenario)[:-1]
    wrong = verdicts(scenario, **{"deck-delivered-and-consistent": {"result": "fail", "session": "ghost"}})
    provider = Provider(
        response(REVIEW, {"verdicts": partial, "remark": "Fine."}),
        response(REVIEW, {"verdicts": wrong, "remark": "Fine."}),
        response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Fine.", "handover": ["price-list"]}),
        response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Fine."}),
    )
    agency = Agency(
        scenario, provider, tmp_path / "skills", workdir=tmp_path, disclosure=scenario.disclosure("all"), deliver="pool"
    )
    agency.prepare()
    assert sorted(agency.released) == sorted(scenario.materials) and agency.withheld == []
    signal = await agency.evaluate({"student": [exchange("Hi", "Hello")]})
    assert signal.satisfied is True
    errors = [json.loads(m["content"])["error"] for m in provider.requests[3]["messages"] if m["role"] == "tool"]
    assert "one verdict per criterion" in errors[0] and "ghost" in errors[1]
    assert "withheld" in errors[2]


async def test_agency_gives_up_within_its_call_budget(tmp_path):
    scenario = Scenario.load(TRAVEL)
    provider = Provider(*(response(REVIEW, {"verdicts": [], "remark": ""}) for _ in range(2)))
    agency = Agency(scenario, provider, tmp_path / "skills", workdir=tmp_path, deliver="pool", max_calls=2)
    with pytest.raises(ExchangeError, match="budget"):
        await agency.evaluate({})
    handed = tmp_path / "handed"
    shutil.copytree(TRAVEL, handed)
    spec = json.loads((handed / "scenario.json").read_text())
    (handed / "scenario.json").write_text(json.dumps({key: value for key, value in spec.items() if key != "plans"}))
    (handed / "materials" / "expert-profile").mkdir()
    (handed / "materials" / "expert-profile" / "profile.md").write_text("# The expert\nRules as written.")
    with pytest.raises(ValueError, match=r"SKILL.md, which \['expert-profile'\] lack"):
        Agency(Scenario.load(handed), provider, tmp_path / "skills", workdir=tmp_path, deliver="pool")
    with pytest.raises(ValueError, match="other materials"):
        Agency(
            scenario.without(["price-list"]),
            provider,
            tmp_path / "skills",
            workdir=tmp_path,
            disclosure=scenario.disclosure("all"),
        )


def test_the_curator_snapshot_does_not_carry_the_scenario_or_the_experiment_docs():
    mounted = [_REPOSITORY / relative for relative in _SOURCE_PATHS]
    hidden = [TRAVEL, _REPOSITORY / "experimental" / "docs", _REPOSITORY / "experimental" / "simulation"]
    assert not any(path.is_relative_to(mount) for path in hidden for mount in mounted)
    assert (_REPOSITORY / "experimental" / "harness" / "curator") in mounted
    showcase = [
        path for path in (_REPOSITORY / "experimental" / "simulation" / "showcase").rglob("*") if path.is_file()
    ]
    assert showcase and not any(path.is_relative_to(mount) for path in showcase for mount in mounted)


def test_a_registered_source_that_climbs_to_the_experimental_package_does_not_widen_the_snapshot(tmp_path):
    """A module under experimental/ resolves its package root to experimental/ itself, which holds the simulation, the
    judges and the showcase; the mount of experimental/harness/curator must keep that root out."""
    repository = tmp_path / "repository"
    for relative, text in {
        "experimental/__init__.py": "",
        "experimental/harness/curator/__init__.py": "",
        "experimental/harness/curator/strategy.py": "def create(): pass\n",
        "experimental/simulation/cases/showcase/transcript/3-round-1.md": "# round 1\n\nquote-sheet-correct: fail\n",
    }.items():
        (repository / relative).parent.mkdir(parents=True, exist_ok=True)
        (repository / relative).write_text(text)
    inspection = SimpleNamespace(
        facts={},
        sources={
            "strategy.create": {
                "path": str(repository / "experimental/harness/curator/strategy.py"),
                "root": str(repository / "experimental"),
                "digest": hashlib.sha256(
                    (repository / "experimental/harness/curator/strategy.py").read_bytes()
                ).hexdigest(),
            }
        },
    )
    exploration = Exploration(
        Config(), inspection, repository=repository, source_paths=("experimental/harness/curator",), root=tmp_path / "x"
    )
    assert set(exploration.mounts) == {"source/experimental/harness/curator"}
    assert not (exploration.root / "source" / "experimental" / "simulation").exists()


def test_the_employees_curator_is_kept_from_the_simulation_its_judges_and_the_scenario(tmp_path):
    repository = tmp_path / "repository"
    files = {
        "raven/helper.py": "VALUE = 1\n",
        "tests/test_raven_helper.py": "def test_value(): pass\n",
        "tests/test_simulation_cards.py": "def test_cards(): pass\n",
        "tests/test_analyst_run.py": "def test_run(): pass\n",
        "tests/test_cards_again.py": "from experimental.simulation.bench.cards import draw\n",
        "tests/test_loop.py": "from experimental.harness.loop.run import run\n",
        "tests/test_named.py": f"SCENARIO = '{TRAVEL.name}'\n",
        "tests/test_scenario_contract.py": "def test_contract(): pass\n",
        "tests/test_assessor_human.py": "def test_human(): pass\n",
        "tests/test_sealed_again.py": "from experimental.harness.scenario.sealed import Sealed\n",
        "tests/test_judging.py": "import experimental.harness.assessor.role\n",
    }
    for relative, text in files.items():
        (repository / relative).parent.mkdir(parents=True, exist_ok=True)
        (repository / relative).write_text(text)
    inspection = SimpleNamespace(facts={}, sources={})
    exploration = Exploration(
        Config(),
        inspection,
        repository=repository,
        source_paths=("raven", "tests"),
        withheld=withheld(Scenario.load(TRAVEL)),
        root=tmp_path / "x",
    )
    assert {path.name for path in (exploration.root / "source" / "tests").iterdir()} == {"test_raven_helper.py"}


@pytest.mark.parametrize("deliver", ["dialog", "pool"])
@pytest.mark.parametrize("without", [(), ("plan-deck-sample",)], ids=["complete", "excluded"])
def test_the_employee_works_from_its_own_copy_of_the_home(tmp_path, deliver, without):
    scenario = Scenario.load(TRAVEL)
    contract = scenario.contract
    scenario = scenario.without(without)
    home = tmp_path / "shared-home"
    for directory in ("skills", "uploads"):
        for name in (*contract.materials, "personal-helper"):
            folder = home / directory / name
            folder.mkdir(parents=True)
            (folder / "personal.txt").write_text("mine")
    (home / "sessions").mkdir()
    (home / "sessions" / "old.jsonl").write_text("{}")
    config = tmp_path / "config.json"
    config.write_text(
        json.dumps(
            {
                "providers": {"deepseek": {"api_key": "k"}},
                "agents": {"defaults": {"model": "deepseek/x", "provider": "deepseek"}},
            }
        )
    )
    root = tmp_path / "run"
    workdir = tmp_path / "work"
    worker = hire(contract, config, workdir=workdir, root=root, home=home, guard=Guard("partner", Sealed.of(["x"])))
    assert worker.withheld == withheld(contract)
    assert isinstance(worker.provider_factory, GuardedFactory) and worker.provider_factory.guard.role == "partner"
    assert hire(contract, config, workdir=tmp_path, root=tmp_path / "plain", home=home).provider_factory is None
    pool = skills(worker)
    assert pool == root / AREA / HOME / "skills" and (pool / "brand-design-guide" / "personal.txt").is_file()
    assert not (root / AREA / HOME / "sessions").exists()
    uploads, shared = root / AREA / HOME / "uploads", workdir / "uploads"
    for name in (*contract.materials, "personal-helper"):
        folder = shared / name
        folder.mkdir(parents=True)
        (folder / "personal.txt").write_text("mine")
    agency = Agency(scenario, Provider(), pool, workdir=workdir, deliver=deliver, uploads=uploads, shared=shared)
    agency.prepare()
    assert agency.released == list(scenario.initial)
    for place in (pool, uploads, shared):
        expected = set(scenario.initial) if (place == pool) == (deliver == "pool") else set()
        assert {name for name in contract.materials if (place / name).exists()} == expected
        for name in expected:
            assert (place / name / "SKILL.md").read_bytes() == (scenario.materials[name] / "SKILL.md").read_bytes()
        assert not any((place / name / "personal.txt").exists() for name in contract.materials)
        assert (place / "personal-helper" / "personal.txt").read_text() == "mine"
    for directory in ("skills", "uploads"):
        assert all((home / directory / name / "personal.txt").read_text() == "mine" for name in contract.materials)


@pytest.mark.parametrize("kind", ["file", "symlink"])
def test_excluded_material_paths_are_removed_without_following_links(tmp_path, kind):
    scenario = Scenario.load(TRAVEL).without(["plan-deck-sample"])
    pool = tmp_path / "skills"
    pool.mkdir()
    retained = tmp_path / "original-material"
    retained.mkdir()
    (retained / "sample.txt").write_text("keep this source")
    stale = pool / "plan-deck-sample"
    if kind == "symlink":
        stale.symlink_to(retained, target_is_directory=True)
    else:
        stale.write_text("stale material")
    Agency(scenario, Provider(), pool, workdir=tmp_path, deliver="pool").prepare()
    assert not stale.exists() and not stale.is_symlink()
    assert (retained / "sample.txt").read_text() == "keep this source"


def test_initial_material_cleanup_failure_prevents_handover(tmp_path, monkeypatch):
    scenario = Scenario.load(TRAVEL)
    pool = tmp_path / "skills"
    stale = pool / "service-sop" / "personal.txt"
    stale.parent.mkdir(parents=True)
    stale.write_text("stale material")
    unlink = os.unlink

    def refuse_stale_file(path, *args, **kwargs):
        if Path(path).name == "personal.txt":
            raise PermissionError("cannot remove stale material")
        return unlink(path, *args, **kwargs)

    monkeypatch.setattr(os, "unlink", refuse_stale_file)
    agency = Agency(scenario, Provider(), pool, workdir=tmp_path, deliver="pool")
    with pytest.raises(PermissionError):
        agency.prepare()
    assert not agency.released
    assert not (pool / "service-sop" / "SKILL.md").exists()


def test_the_employee_houses_its_external_subagents_homes_beside_its_own_harness(tmp_path, monkeypatch):
    from experimental.cultivation import employee
    from raven.config.schema import ThirdPartyAcpSubagentConfig

    def row(name, enabled=True):
        return ThirdPartyAcpSubagentConfig(name=name, kind="acp", command="run", enabled=enabled, env={"A": "1"})

    monkeypatch.setattr(employee, "discover_product_rows", lambda: [row("Raven-PPT"), row("Raven-Research", False)])
    monkeypatch.setenv("PYTHONTZPATH", "/zones")
    config = tmp_path / "config.json"
    config.write_text(
        json.dumps(
            {
                "providers": {"deepseek": {"api_key": "k"}},
                "agents": {"defaults": {"model": "deepseek/x", "provider": "deepseek"}},
            }
        )
    )
    worker = hire(Scenario.load(TRAVEL).contract, config, workdir=tmp_path, root=tmp_path / "run")
    rows = {item.name: item for item in worker.baseline.config.subagents.agents}
    housed = tmp_path / "run" / AREA / HOME / "subagents" / "Raven-PPT"
    assert rows["Raven-PPT"].env == {"A": "1", "PYTHONTZPATH": "/zones", "PPT_ACP_HOME": str(housed)}
    assert housed.is_dir() and "Raven-Research" not in rows


def test_the_housed_subagents_run_on_a_model_of_their_own_through_the_employees_provider(tmp_path, monkeypatch):
    from experimental.cultivation import employee
    from raven.config.loader import load_config
    from raven.config.schema import ThirdPartyAcpSubagentConfig
    from raven.providers.endpoints import provider_endpoints

    def row(name):
        return ThirdPartyAcpSubagentConfig(name=name, kind="acp", command=f"py /agents/{name}/run.py", enabled=True)

    folder = tmp_path / "product"
    folder.mkdir()
    shipped = {
        "agents": {
            "defaults": {
                "model": "vendor/other",
                "provider": "ppt",
                "maxTokens": 9,
                "reasoningEffort": "max",
                "temperature": 0.9,
                "llmCallTimeout": 900,
            }
        },
        "providers": {"ppt": {"apiBase": "https://gateway.example/v1"}},
        "acp": {
            "modes": {
                "medium": {"name": "Medium", "reasoningEffort": "low"},
                "high": {"name": "High", "reasoningEffort": "high"},
            },
            "defaultMode": "high",
        },
        "tools": {"disabledTools": ["spawn"]},
        "plugins": {"config": {"research-flow": {"verify": {"reasoningEffort": "low", "maxTokens": 16384}}}},
    }
    (folder / "config.json").write_text(json.dumps(shipped))
    monkeypatch.setattr(
        employee, "discover_product_rows", lambda: [row("Raven-PPT"), row("Raven-Research"), row("Raven-Oncall")]
    )
    monkeypatch.setattr(employee, "product_folder", lambda name: folder)
    connection = {
        "apiKey": "ds-key",
        "apiBase": "https://employee-gateway.example/v1",
        "extraHeaders": {"X-Route": "employee"},
        "protocol": "chat",
        "modelProtocols": {"deepseek/deepseek-flash": "chat"},
    }
    settings = {
        "providers": {"deepseek": connection},
        "permissions": {"mode": "full", "judgeTimeoutSeconds": 60},
        "acp": {"defaultMode": "medium"},
        "tools": {"restrictToWorkspace": True},
        "agents": {
            "defaults": {
                "model": "deepseek/deepseek-flash",
                "provider": "deepseek",
                "reasoningEffort": "low",
                "contextWindowTokens": 1048576,
            }
        },
    }
    config = tmp_path / "config.json"
    config.write_text(json.dumps(settings))
    worker = hire(
        Scenario.load(TRAVEL).contract,
        config,
        workdir=tmp_path,
        root=tmp_path / "run",
        subagent_model="deepseek/deepseek-flash",
    )
    rows = {item.name: item for item in worker.baseline.config.subagents.agents}
    for name, secret in (
        ("Raven-PPT", "PPT_API_KEY"),
        ("Raven-Research", "RESEARCH_API_KEY"),
        ("Raven-Oncall", "ONCALL_API_KEY"),
    ):
        copy = tmp_path / "run" / AREA / "deployment" / f"{name.lower()}.json"
        written = json.loads(copy.read_text())
        assert rows[name].env[secret] == "ds-key" and rows[name].command.endswith(f"--config {copy}")
        assert written["agents"]["defaults"] == {
            "model": "deepseek/deepseek-flash",
            "provider": "deepseek",
            "maxTokens": 9,
            "reasoningEffort": "low",
            "contextWindowTokens": 1048576,
            "temperature": 0.9,
            "llmCallTimeout": 900,
        }
        assert written["providers"]["deepseek"] == {**connection, "models": ["deepseek/deepseek-flash"]}
        assert copy.stat().st_mode & 0o777 == 0o600
        restored = load_config(copy)
        assert restored.get_api_base() == connection["apiBase"]
        assert restored.get_provider().extra_headers == connection["extraHeaders"]
        assert provider_endpoints(restored.get_provider()) == provider_endpoints(
            worker.baseline.config.providers.deepseek
        )
        assert restored.agents.defaults.temperature == 0.9 and restored.agents.defaults.llm_call_timeout == 900
        assert written["permissions"] == {"mode": "full", "judgeTimeoutSeconds": 60.0}
        # The shipped high tier would run every deck call at high effort over the pinned low one.
        assert written["acp"]["defaultMode"] == "medium"
        assert written["tools"]["disabledTools"] == (
            ["spawn", "ask_user"] if name == "Raven-Oncall" else ["spawn", "ask_user", "exec"]
        )
        # The employee's workspace fence goes along: a product reads only its home and the session's workdir.
        assert written["tools"]["restrictToWorkspace"] is True
        # A thinking DeepSeek model refuses the verify gate's forced verdict call, and the gate then passes every draft.
        verify = written["plugins"]["config"]["research-flow"]["verify"]
        assert verify == {"reasoningEffort": "none", "maxTokens": 16384}
    assert "PPT_ACP_HOME" in rows["Raven-PPT"].env
    assert worker.baseline.config.providers.deepseek.models == ["deepseek/deepseek-flash"]
    settings["agents"]["defaults"].pop("reasoningEffort")
    config.write_text(json.dumps(settings))
    with pytest.raises(ValueError, match="reasoning effort"):
        hire(
            Scenario.load(TRAVEL).contract,
            config,
            workdir=tmp_path,
            root=tmp_path / "again",
            subagent_model="deepseek/x",
        )


def test_a_raven_child_on_a_gateway_model_takes_the_employee_provider(monkeypatch):
    from raven.config.product_render import inherit_llm
    from raven.config.schema import Config

    config = Config.model_validate(
        {
            "providers": {"openrouter": {"apiKey": "or-key"}},
            "agents": {"defaults": {"model": "deepseek/deepseek-v4-flash", "provider": "openrouter"}},
        }
    )
    monkeypatch.setenv("RAVEN_PARENT_MODEL", "deepseek/deepseek-v4-flash")
    monkeypatch.delenv("RAVEN_PARENT_PROVIDER", raising=False)
    host = config.model_dump(mode="json", by_alias=True)
    assert "provider=deepseek " in inherit_llm({}, host)
    listed(config, config.agents.defaults.model, "deepseek/deepseek-v4-flash", None)
    assert config.providers.openrouter.models == ["deepseek/deepseek-v4-flash"]
    assert "provider=openrouter " in inherit_llm({}, config.model_dump(mode="json", by_alias=True))


async def test_the_agency_reads_the_last_version_of_each_delivered_file(tmp_path, monkeypatch):
    from experimental.harness.participation.views import opened, transcript
    from experimental.harness.raven_adapter.worker import Execution

    if shutil.which("soffice") is None:
        from pptx import Presentation

        from experimental.harness.participation import files
        from experimental.harness.participation.render import cached

        def draw(file, out, width):
            for number in range(1, len(Presentation(file).slides) + 1):
                (out / f"page-{number:02d}.png").write_bytes(b"png")

        monkeypatch.setattr(files, "cached", lambda file, cache, width: cached(file, cache, width, draw))
    scenario = Scenario.load(TRAVEL)
    first, second = tmp_path / "t1" / "trip.html", tmp_path / "t2" / "trip.html"
    final = '<style>h1{color:red}</style><h1>final</h1><img src="data:image/png;base64,' + "A" * 300 + '">'
    for path, text in ((first, "<h1>draft</h1>"), (second, final)):
        path.parent.mkdir()
        path.write_text(text)
    deck = tmp_path / "t3" / "plan.pptx"
    deck.parent.mkdir()
    shutil.copy(scenario.materials["plan-deck-template"] / "template.pptx", deck)
    records = [{"kind": "runner.event", "event_type": "Text", "event": {"content": "Here it is."}}]
    exchanges = [
        Exchange("Plan it", Execution("t1", [], records, {}, "a", (str(first),))),
        Exchange("Fix it", Execution("t2", [], records, {}, "a", (str(second),))),
        Exchange("The deck?", Execution("t3", [], records, {}, "a", (str(deck),))),
    ]
    monkeypatch.setattr("experimental.cultivation.counterpart.owner.vision_verdict", lambda model: model == "sees")
    provider = Provider(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Good page."}))
    await Agency(scenario, provider, tmp_path / "skills", workdir=tmp_path, deliver="pool", model="sees").evaluate(
        {"student": exchanges}
    )
    packet = json.loads(provider.requests[0]["messages"][1]["content"])
    delivered = packet["deliverables"]["student"]
    assert packet["references"]["student"]["deck"]["template_placeholders_left"]
    assert delivered["trip.html"] == "final"
    assert (
        delivered["plan.pptx"].startswith("19 slides, aspect 1.78") and "[style] background" in delivered["plan.pptx"]
    )
    assert packet["pictures"] == {"student": {"plan.pptx": 19}}
    pictures = provider.requests[0]["messages"][2]["content"]
    images = [part for part in pictures if part["type"] == "image_url"]
    assert len(images) == 19 and images[0]["image_url"]["url"].startswith("data:image/png;base64,")
    assert (deck.parent / "plan.pptx.thumbs" / "page-01.png").is_file()
    blind = Provider(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Read it."}))
    await Agency(scenario, blind, tmp_path / "skills", workdir=tmp_path, deliver="pool", model="text").evaluate(
        {"student": exchanges}
    )
    unseen = json.loads(blind.requests[0]["messages"][1]["content"])["pictures"]
    assert unseen["decks"] == {"student": {"plan.pptx": 19}} and "not shown" in unseen["not_shown"]
    assert not any(isinstance(message.get("content"), list) for message in blind.requests[0]["messages"])
    assert transcript(exchanges)[0]["delivered"] == ["trip.html"] and "stopped" not in transcript(exchanges)[0]
    late = Execution("t9", [], [], {"timed_out": True, "timeout": 600, "explicit_reply": False})
    assert transcript([Exchange("Hello?", late)])[0]["stopped"] == "no reply within 600 seconds; the turn was stopped"
    seen = opened(exchanges)
    assert seen["trip.html"] == "final" and "--- slide 1\n" in seen["plan.pptx"] and "[style]" not in seen["plan.pptx"]


def test_the_agency_reads_what_the_research_colleague_reported_in_each_playbook_run(tmp_path):
    from experimental.harness.raven_adapter.worker import Execution
    from experimental.simulation.cases.travel_agency.owner import research

    found, drafted = tmp_path / "research.out.md", tmp_path / "brief.out.md"
    found.write_text("G7311 Shanghai to Huangshan, 3h (source: 12306)")
    drafted.write_text("brief")
    manifest = {
        "files": [
            {"node": "p-research", "subagent": "Raven-Research", "output_file": str(found)},
            {"node": "p-brief", "subagent": "Raven", "output_file": str(drafted)},
        ]
    }
    records = [{"kind": "dag.progress", "name": "dag_run_completed", "payload": {"manifest": manifest}}]
    exchanges = [Exchange("Deck please", Execution("t1", [], records, {}, "a"))]
    assert research(exchanges) == {"p-research": "G7311 Shanghai to Huangshan, 3h (source: 12306)"}


async def test_the_owner_gets_the_played_card_and_its_figures_and_still_decides_every_verdict(tmp_path):
    from experimental.simulation.bench.cards import Drawn, Trip

    scenario = Scenario.load(TRAVEL)
    sheet = scenario.text("service-sop").split("```")[1].strip()
    guest = Trip("guest", "origin", "somewhere", (12, 2), (12, 6), 4, 4, (), (), 20_000, "139 0571 6628")
    card = Drawn("student", "card text", guest)
    provider = Provider(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Fine."}))
    agency = Agency(
        scenario,
        provider,
        tmp_path / "skills",
        workdir=tmp_path,
        disclosure=scenario.disclosure("all"),
        deliver="pool",
        cards=lambda: {"student": card},
        records=tmp_path / "records",
    )
    agency.prepare()
    signal = await agency.evaluate({"student": [exchange("Quote please", sheet)]})
    packet = json.loads(provider.requests[0]["messages"][1]["content"])
    assert packet["cards"] == {"student": "card text"}
    comfort = list(packet["references"]["student"]["prices"]["by_product"].values())[1]
    assert comfort["party_total"] == 12_800 and "deck" not in packet["references"]["student"]
    assert signal.satisfied is True and len(provider.requests) == 1
    rows = [json.loads(line) for line in (tmp_path / "records" / "references.jsonl").read_text().splitlines()]
    assert rows[0]["drill"] == "student" and rows[0]["trip"]["adults"] == 4 and rows[0]["card"] == "card text"
    assert rows[0]["references"]["prices"]["persons"] == 4


def test_a_scenario_without_a_material_loses_it_from_its_materials_opening_set_and_plans():
    scenario = Scenario.load(TRAVEL)
    control = scenario.without(["plan-deck-sample"])
    assert "plan-deck-sample" in scenario.materials and "plan-deck-sample" not in control.materials
    assert "plan-deck-sample" not in control.initial and len(control.initial) == len(scenario.initial) - 1
    assert all("plan-deck-sample" not in step for plan in control.plans.values() for step in plan.steps)
    assert control.disclosure("by-stage").steps == control.plans["by-stage"].steps
    assert control.disclosure("by-stage").labels == scenario.plans["by-stage"].labels
    with pytest.raises(ValueError, match="brochure"):
        scenario.without(["brochure"])


def test_the_employees_workdir_is_new_and_never_the_repository(tmp_path):
    from experimental.cultivation.employee import workplace

    repository = Path(__file__).resolve().parents[1]
    for bad in (repository, repository / "experimental", repository.parent):
        with pytest.raises(ValueError, match="repository"):
            workplace(bad)
    (tmp_path / "used").mkdir()
    (tmp_path / "used" / "notes.md").write_text("x")
    with pytest.raises(ValueError, match="new or empty"):
        workplace(tmp_path / "used")
    assert workplace(tmp_path / "fresh").is_dir() and not any(workplace(None).iterdir())


def test_each_child_reaches_only_its_own_home_and_what_its_parent_hands_it(tmp_path, monkeypatch):
    """The children `hire` wires are prepared as the worker would prepare them on start."""
    from experimental.cultivation import employee

    monkeypatch.setattr(employee, "discover_product_rows", lambda: [])
    config = tmp_path / "config.json"
    config.write_text(
        json.dumps(
            {
                "providers": {"deepseek": {"api_key": "k"}},
                "agents": {"defaults": {"model": "deepseek/x", "provider": "deepseek"}},
            }
        )
    )
    worker = hire(Scenario.load(TRAVEL).contract, config, workdir=tmp_path / "work", root=tmp_path / "run")
    children = worker._prepare_children(worker.baseline)
    assert set(children) == {"Raven"}
    child = children["Raven"].baseline
    assert child.config.workspace_path == worker.baseline.config.workspace_path / "subagents" / "Raven"
    assert child.file_roots == (worker.baseline.workdir,) == (tmp_path / "work",)
    assert child.read_roots == (worker.baseline.config.workspace_path,) == (tmp_path / "run" / AREA / HOME,)
    assert child.config.tools.restrict_to_workspace is True
    assert worker.baseline.file_roots == () and worker.baseline.read_roots == ()


async def test_the_owner_hands_over_a_candidate_it_never_confirmed_as_files_and_never_voices_it(tmp_path):
    from experimental.harness.scenario import load
    from tests.research_helpers import make_scenario

    CONFIRMED, CANDIDATES = "adopted-method", "candidate-method"
    root = make_scenario(tmp_path / "shop")
    contract = json.loads((root / "contract.json").read_text())
    spec = json.loads((root / "scenario.json").read_text())
    for name, text in ((CONFIRMED, "State the total price."), (CANDIDATES, "Never offer a discount.")):
        folder = root / "materials" / name
        folder.mkdir()
        (folder / "SKILL.md").write_text(f"---\nname: {name}\ndescription: Supplied method\n---\n{text}\n")
        contract["materials"][name] = {"kind": "norm"}
        spec["plans"]["two"]["steps"][1].append(name)
    (root / "contract.json").write_text(json.dumps(contract))
    (root / "scenario.json").write_text(json.dumps(spec))
    (root / "provenance.json").write_text(
        json.dumps(
            {
                "materials": {
                    CONFIRMED: {"origin": "given", "confirmed": True},
                    CANDIDATES: {"origin": "given", "confirmed": False},
                }
            }
        )
    )
    scenario = Scenario.of(load(root))
    provider = Provider(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Fine so far."}))
    uploads = tmp_path / "home" / "uploads"
    agency = Agency(
        scenario,
        provider,
        tmp_path / "home" / "skills",
        workdir=tmp_path,
        uploads=uploads,
        disclosure=scenario.disclosure("two"),
    )
    agency.prepare()
    signal = await agency.evaluate({"walk-in": [exchange("Tea for two?", "Six.")]})
    handed = {item.name for item in signal.attachments}
    assert {CONFIRMED, CANDIDATES} <= handed and (uploads / CANDIDATES / "SKILL.md").is_file()
    assert scenario.document(CONFIRMED).strip() in signal.text
    assert scenario.document(CANDIDATES).strip() not in signal.text and "Never offer a discount." not in signal.text


def test_held_out_cards_are_played_apart_from_the_cards_played_for_feedback():
    from experimental.simulation.__main__ import drills

    personas = Scenario.load(TRAVEL).personas
    names = [persona.name for persona in personas]
    played, held = drills(personas, None, [names[0]])
    assert [persona.name for persona in held] == [names[0]] and [persona.name for persona in played] == names[1:]
    played, held = drills(personas, names[1:2], None)
    assert [persona.name for persona in played] == names[1:2] and held == []
    with pytest.raises(ValueError, match="not also played for feedback"):
        drills(personas, names[:1], names[:1])
    with pytest.raises(ValueError, match="no drill cards named"):
        drills(personas, None, ["nobody"])
    with pytest.raises(ValueError, match="leaves nothing to play"):
        drills(personas, None, names)


def test_cards_and_exam_named_on_the_command_line_each_replace_their_part_of_the_schedule():
    from experimental.harness.scenario.contract import Schedule
    from experimental.simulation.__main__ import exam_rounds, scheduled

    personas = Scenario.load(TRAVEL).personas
    names = lambda cards: [card.name for card in cards]  # noqa: E731
    schedule = Schedule(rounds=(("family",), ("student",)), holdout=("returning",), holdout_rounds=(1, 2))
    rounds, held = scheduled(personas, schedule)
    assert [names(step) for step in rounds] == [["family"], ["student"]] and names(held) == ["returning"]
    rounds, held = scheduled(personas, schedule, cards=["premium"])
    assert [names(step) for step in rounds] == [["premium"]] and names(held) == ["returning"]
    rounds, held = scheduled(personas, schedule, holdout=["professional"])
    assert [names(step) for step in rounds] == [["family"], ["student"]] and names(held) == ["professional"]
    with pytest.raises(ValueError, match="not also played"):
        scheduled(personas, schedule, holdout=["family"])
    rounds, held = scheduled(personas, Schedule(holdout=("returning",)))
    assert "returning" not in names(rounds[0]) and names(held) == ["returning"]
    assert exam_rounds(schedule, None, 3) == frozenset({1, 2}) and exam_rounds(schedule, [3], 3) == frozenset({3})
    assert exam_rounds(Schedule(), None, 3) is None
    with pytest.raises(ValueError, match="rounds \\[4\\]"):
        exam_rounds(schedule, [1, 4], 3)
    assert exam_rounds(schedule, [0, 3], 3) == frozenset({0, 3})
    with pytest.raises(ValueError, match="rounds \\[-1\\]"):
        exam_rounds(schedule, [-1], 3)


def test_a_case_is_judged_by_the_owner_of_the_bundled_case_it_is_named_after(tmp_path):
    from experimental.simulation.__main__ import owner_of
    from experimental.simulation.cases.brand_strategist.owner import BrandOwner

    assert owner_of(Scenario.load(BUNDLED / "brand_strategist")) is BrandOwner
    assert owner_of(Scenario.load(TRAVEL)) is Agency
    shutil.copytree(BUNDLED / "brand_strategist", tmp_path / "brand_strategist")
    assert owner_of(Scenario.load(tmp_path / "brand_strategist")) is BrandOwner
    shutil.copytree(BUNDLED / "travel_agency", tmp_path / "elsewhere")
    with pytest.raises(ValueError, match="not a bundled case"):
        owner_of(Scenario.load(tmp_path / "elsewhere"))
    (tmp_path / "elsewhere" / "owner.py").unlink()
    assert owner_of(Scenario.load(tmp_path / "elsewhere")) is Agency


def test_a_played_card_carries_the_expected_answer_its_case_labels(tmp_path):
    import json
    import random

    from experimental.harness.scenario.contract import frontmatter

    root = tmp_path / "travel_agency"
    shutil.copytree(TRAVEL, root)
    card = root / "personas" / "student.md"
    head, body = frontmatter(card.read_text())
    card.write_text("---\n" + json.dumps({**head, "expected": "Quote the value package only."}) + "\n---\n" + body)
    (student,) = [persona for persona in Scenario.load(root).personas if persona.name == "student"]
    assert student.play(random.Random(1)).expected == "Quote the value package only."


async def test_the_owner_judges_a_labelled_card_with_its_expected_answer(tmp_path):
    from experimental.simulation.bench.cards import Drawn

    scenario = Scenario.load(TRAVEL)
    provider = Provider(response(REVIEW, {"verdicts": verdicts(scenario), "remark": "Fine."}))
    played = {"student": Drawn("student", "You want a cheap trip.", None, "Only the value package fits the budget.")}
    agency = Agency(
        scenario,
        provider,
        tmp_path / "skills",
        workdir=tmp_path,
        deliver="pool",
        cards=lambda: played,
        handouts=lambda sessions: {"student": {"proof.md": "The purchased voucher lasts 90 days."}},
        records=tmp_path / "run",
        model="sim",
    )
    agency.prepare()
    await agency.evaluate({"student": [exchange("Cheap trip?", "The value package fits.")]})
    packet = json.loads(provider.requests[0]["messages"][1]["content"])
    assert packet["expected"] == {"student": "Only the value package fits the budget."}
    assert packet["handouts"]["student"]["proof.md"] == "The purchased voucher lasts 90 days."
    assert "price-list" in packet["material_files"]
    (row,) = [json.loads(line) for line in (tmp_path / "run" / "references.jsonl").read_text().splitlines()]
    assert row["expected"] == "Only the value package fits the budget."


def test_evaluator_sources_include_visible_facts_and_nested_files_without_releasing_hidden_materials(tmp_path):
    from dataclasses import replace

    folder = tmp_path / "travel_agency"
    shutil.copytree(TRAVEL, folder)
    nested = folder / "materials/price-list/details.md"
    nested.write_text("Adopted factual detail.")
    scenario = Scenario.load(folder)
    material = scenario.contract.materials["price-list"].model_copy(update={"visibility": frozenset({"counterpart"})})
    scenario = replace(
        scenario,
        contract=scenario.contract.model_copy(
            update={
                "materials": {**scenario.contract.materials, "price-list": material},
            }
        ),
    )
    assert scenario.evidence_materials(["price-list"], role="analyst") == {}
    evidence = scenario.evidence_materials(["price-list"], role="counterpart")
    assert "price-list" in evidence
    assert evidence["price-list/details.md"] == "Adopted factual detail."
    assert "price-list" not in scenario.standing_norms(["price-list"])
    assert scenario.evidence_materials(["service-sop"], role="counterpart").keys() == {"service-sop"}


def test_judge_receives_each_cards_handouts_without_private_role_or_answer(tmp_path):
    from experimental.simulation.bench.scenario import handed_evidence

    one, other = tmp_path / "one.md", tmp_path / "other.md"
    one.write_text("The voucher is valid for 90 days.")
    other.write_text("An unrelated customer's file.")
    conversants = {
        "played": SimpleNamespace(handed=(one,), card=SimpleNamespace(text="Private role", expected="Private answer")),
        "unplayed": SimpleNamespace(handed=(other,)),
    }
    evidence = handed_evidence({"played": []}, conversants)
    assert evidence == {"played": {"handouts": {"one.md": "The voucher is valid for 90 days."}}}


async def test_owner_requires_a_declared_condition_and_evidence_for_not_applicable(tmp_path):
    from dataclasses import replace

    scenario, agency = owner(tmp_path, Provider())
    review = scorecard(scenario)
    review["verdicts"][0].update(result="not_applicable", note="All requested dates precede the effective date.")
    names = list(handles(scenario.criteria).values())
    sessions = {"student": [exchange("Use the earlier date.", "Understood.")]}
    with pytest.raises(ValueError, match="unconditional"):
        agency._parse(sessions, review, names)
    first = scenario.criteria[0].model_copy(
        update={"applies_when": "The requested date is after the policy takes effect."}
    )
    agency.scenario = replace(scenario, criteria=(first, *scenario.criteria[1:]))
    review["verdicts"][0]["note"] = " "
    with pytest.raises(ValueError, match="evidenced explanation"):
        agency._parse(sessions, review, names)
    review["verdicts"][0]["note"] = "All requested dates precede the effective date."
    agency.provider = Provider(response(REVIEW, review))
    signal = await agency.evaluate(sessions)
    assert signal.satisfied is True
    kept = next((tmp_path / "run" / "analysis").glob("*.json"))
    measured = json.loads(kept.read_text())["scorecard"]["items"][0]
    assert measured["result"] == "not_applicable" and measured["note"] == review["verdicts"][0]["note"]
    agency.scenario = replace(scenario, criteria=(first,))
    agency.provider = Provider(response(REVIEW, {**review, "verdicts": review["verdicts"][:1]}))
    assert (await agency.evaluate(sessions)).satisfied is None


async def test_private_owner_evidence_is_retained_exactly_without_entering_feedback(tmp_path):
    scenario = Scenario.load(TRAVEL)
    provider = Provider(response(REVIEW, scorecard(scenario, remark="Please improve the process.")))
    scenario, agency = owner(tmp_path, provider)
    body = "Internal filed note. " + "Detail retained beyond two hundred characters. " * 20
    agency._filed = lambda name: {"workdir/department/review.md": body}
    agency._provided_handouts = lambda sessions: {"student": {"customer.md": "Supplied customer source."}}
    sessions = {"student": [exchange("Please check the request.", "I filed it.")]}
    spoken = await agency.evaluate(sessions)
    kept = json.loads(next((tmp_path / "run/analysis").glob("*.json")).read_text())
    submitted = json.loads(provider.requests[0]["messages"][1]["content"])
    assert kept["evidence"] == submitted
    assert kept["evidence"]["filed"]["student"]["workdir/department/review.md"] == body
    assert kept["execution_ids"] == {"student": [sessions["student"][0].execution.turn_id]}
    assert kept["rubric_version"] and kept["assessment_id"]
    assert kept["criterion_handles"] == handles(scenario.criteria)
    assert {row["id"] for row in kept["criteria"]} == {criterion.id for criterion in scenario.criteria}
    assert "Internal filed note" not in spoken.text and "Supplied customer source" not in spoken.text
    assert spoken.items == () and spoken.metrics == {}


def test_filed_evidence_never_reads_escaped_current_or_baseline_links(tmp_path, monkeypatch):
    from experimental.cultivation.counterpart.owner import written

    current, before = tmp_path / "current", tmp_path / "before"
    current.mkdir()
    before.mkdir()
    outside = tmp_path / "outside.md"
    outside.write_text("Outside content must not enter the grader packet.")
    (current / "escape.md").symlink_to(outside)
    (current / "comparison.md").write_text("A current file with an unverifiable baseline.")
    (before / "comparison.md").symlink_to(outside)
    (current / "valid.md").write_text("A valid in-scope record.")
    (current / "alias.md").symlink_to(current / "valid.md")
    original = Path.read_bytes

    def checked(path):
        assert path.resolve() != outside.resolve()
        return original(path)

    monkeypatch.setattr(Path, "read_bytes", checked)
    files = written(current, before, label="workdir")
    assert files == {"workdir/alias.md": "A valid in-scope record.", "workdir/valid.md": "A valid in-scope record."}
    aliased_root = tmp_path / "root-link"
    aliased_root.symlink_to(current, target_is_directory=True)
    assert written(aliased_root, before, label="workdir") == {}
